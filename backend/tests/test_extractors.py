"""
Extractor tests — every reader from the Phase 4 brief, tested against REAL
files generated in-test (fpdf2 makes genuine PDFs, python-docx real .docx, …).
"""

from __future__ import annotations

from pathlib import Path

import pytest

from services.extractors import ExtractionError, extract_text


def test_text_markdown_and_code(tmp_path: Path):
    f = tmp_path / "notes.md"
    f.write_text("# नोट्स\n\n- markdown works\n- `code` too\n", encoding="utf-8")
    assert "नोट्स" in extract_text(f, "text", f.name)
    py = tmp_path / "app.py"
    py.write_text("print('vednix')\n", encoding="utf-8")
    assert "vednix" in extract_text(py, "text", py.name)


def test_csv_extraction(tmp_path: Path):
    f = tmp_path / "portfolio.csv"
    f.write_text("symbol,qty,price\nRELIANCE,10,2450\nTCS,5,3900\n", encoding="utf-8")
    out = extract_text(f, "sheet", f.name)
    assert "symbol | qty | price" in out
    assert "RELIANCE | 10 | 2450" in out


def test_xlsx_extraction(tmp_path: Path):
    import openpyxl

    f = tmp_path / "budget.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "January"
    ws.append(["Item", "Cost"])
    ws.append(["Server", 4200])
    wb.save(f)
    out = extract_text(f, "sheet", f.name)
    assert "Sheet: January" in out and "Server | 4200" in out


def test_docx_extraction(tmp_path: Path):
    import docx

    f = tmp_path / "report.docx"
    document = docx.Document()
    document.add_paragraph("Quarterly results were strong.")
    document.add_paragraph("भारत में ग्रोथ अच्छी रही।")
    document.save(str(f))
    out = extract_text(f, "document", f.name)
    assert "Quarterly results were strong." in out
    assert "भारत में ग्रोथ" in out


def test_pptx_extraction(tmp_path: Path):
    from pptx import Presentation
    from pptx.util import Inches

    f = tmp_path / "deck.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Vednix Roadmap"
    left = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(1))
    left.text_frame.text = "Phase 4: files and vision"
    deck.save(str(f))
    out = extract_text(f, "slides", f.name)
    assert "Slide 1" in out and "Vednix Roadmap" in out and "Phase 4" in out


def test_pdf_extraction_real_file(tmp_path: Path):
    from fpdf import FPDF

    f = tmp_path / "paper.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=14)
    pdf.cell(text="Vednix reads PDFs offline.")
    pdf.output(str(f))
    out = extract_text(f, "document", f.name)
    assert "Vednix reads PDFs offline." in out


def test_image_raises_guidance(tmp_path: Path):
    f = tmp_path / "pic.png"
    f.write_bytes(b"\x89PNG\r\n\x1a\n")  # header bytes are enough for kind routing
    with pytest.raises(ExtractionError, match="vision model"):
        extract_text(f, "image", f.name)


def test_corrupt_file_fails_cleanly(tmp_path: Path):
    f = tmp_path / "broken.xlsx"
    f.write_bytes(b"definitely not a spreadsheet")
    with pytest.raises(ExtractionError):
        extract_text(f, "sheet", f.name)
